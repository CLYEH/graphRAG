"""Source connectors: external data → uniform document payloads (DESIGN §5 step 1, C2).

Two §2 source families feed the pipeline: free-text documents and structured
(tabular) data. Both normalize to :class:`DocumentPayload` here, so the rest
of ingest — hashing, dedup, persistence — is connector-agnostic. Structured
rows become one payload each, carrying ``{"table", "pk"}`` metadata because
§27.2's ``row`` source refs must later cite ``table + pk``; a row that loses
its pk at ingest could never be cited at query time.

``content_hash`` is THE document identity: §18 fixes ``item_ref =
content_hash`` for documents (rerun line-up), and §5's skip/rerun idempotency
compares it. It hashes the raw text only — metadata/source_uri may change
(file moved, re-exported) without making the content a "new" document.

Connectors are synchronous and yield lazily: sources can be large, and the
async persistence boundary is `core.ingest.documents`, not file IO.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from core.metadata.schema import MetadataSchema, MetadataValidationError, build_envelope
from core.stores.tables import STRUCTURED_MIME

#: Free-text suffixes the document connector accepts, mapped to their mime
#: type. Extend as real projects need more formats (PDF etc. arrive with
#: their own extraction dependencies — deliberately not in C2).
TEXT_SUFFIXES: dict[str, str] = {
    ".txt": "text/plain",
    ".md": "text/markdown",
}


@dataclass(frozen=True)
class DocumentPayload:
    """One ingestable unit, whatever the source (§5 step 1)."""

    source_uri: str
    raw: str
    mime: str
    metadata: dict[str, Any] = field(default_factory=dict)


def content_hash(raw: str) -> str:
    """The stable document identity (§18 item_ref; §5 idempotency key)."""
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def read_text_documents(
    root: Path,
    metadata_by_filename: Mapping[str, dict[str, Any]] | None = None,
    *,
    metadata_schema: MetadataSchema | None = None,
) -> Iterator[DocumentPayload]:
    """Yield one payload per accepted text file under ``root`` (recursive).

    Deterministic order (sorted paths) so two runs over the same tree produce
    the same payload sequence — rerun line-up starts at the connector. Files
    with unaccepted suffixes are silently not SELECTED (they are out of scope
    by definition, not failed items); an accepted file that cannot be decoded
    raises — a corrupt source is a loud failure, not a skipped item.

    ``metadata_by_filename`` carries the DR-010 metadata envelope captured at
    upload time, keyed by the file's stored (on-disk) name — the managed-source
    stash the ingest connector threads onto ``documents.metadata`` (UXC1b
    capture → persist). When it is PRESENT (not None — even an empty map) the
    source is a MANAGED one whose registered file list is AUTHORITATIVE in BOTH
    directions: it is the ITERATION SOURCE, so (a) an on-disk file NOT registered —
    an orphan left by a failed / rolled-back upload — is never ingested with
    fallback metadata, and (b) a registered file MISSING from disk fails LOUDLY
    (the SoR says the upload was accepted; silently ingesting fewer documents would
    corrupt results). Presence, not truthiness, is the managed signal: an empty map
    ingests NOTHING (an empty authoritative list), never a directory scan — a
    managed source must not silently degrade to reading unregistered files. The
    keys are UNTRUSTED (a source's managed-file stash is stored as-is), so each
    is validated to a bare in-root filename before it is joined to ``root`` — a
    name with a path separator / dot segment / absolute path is refused, never
    read. An ABSENT (None) stash is a plain directory source: every accepted file
    under the tree is read, each falling back to the connector-derived
    ``{"filename": ...}`` — the original behavior for a non-upload source or a file
    placed on disk directly.

    ``metadata_schema`` fences the SIDECAR path (MCP10 / Codex #130): sidecar
    ``context.attributes`` are validated against the project ``metadata_schema``
    — the SAME fence the upload boundary applies at capture — so an undeclared /
    wrong-typed / missing-required attribute fails the build loud instead of
    persisting values that typed document filters would silently disagree with.
    ``None`` folds to the EMPTY schema (fail-closed: every attribute is then
    undeclared — the schema is load-bearing even when empty), so a caller that
    forgets to thread the project schema breaks a legitimate sidecar build
    loudly rather than silently skipping the fence. Files WITHOUT a sidecar
    never enter the envelope surface, so nothing is validated for them; the
    managed path's envelopes were validated at upload capture.
    """
    if not root.is_dir():
        raise NotADirectoryError(f"document source root {root} is not a directory")
    if metadata_by_filename is not None:
        base = root.resolve()
        # Managed source: iterate the REGISTERED list (sorted), not the directory.
        for name in sorted(metadata_by_filename):
            # The registered names are UNTRUSTED: a text source's managed-file stash
            # is stored as-is by the sources API, so a key like '../other/secret.md'
            # or an absolute '/etc/passwd' would make `root / name` read OUTSIDE the
            # source root (a build ingesting unrelated local files). Require a bare
            # in-root filename — no path separators, '.', '..', or absolute paths —
            # then confirm the RESOLVED path is a DIRECT child of root (a symlink
            # backstop). An unsafe name is a config error at the door, never a read
            # outside the root. (Uploads mint UUID-hex stored names, so this only
            # ever rejects a hand-registered malicious/malformed source.)
            if "\\" in name or name in {"", ".", ".."} or name != Path(name).name:
                raise ValueError(
                    f"registered upload file name {name!r} is not a bare in-root filename "
                    "(no path separators, '.', '..', or absolute paths) — it would read "
                    "outside the source root"
                )
            path = root / name
            if path.resolve().parent != base:
                raise ValueError(
                    f"registered upload file name {name!r} resolves outside the source "
                    f"root {root} (a symlink escaping the corpus)"
                )
            suffix = path.suffix.lower()
            if not path.is_file():
                raise FileNotFoundError(
                    f"registered upload file {name!r} is missing from {root} — the "
                    "managed source lists it but it is not on disk (a lost write); "
                    "ingesting fewer documents than the SoR accepted would corrupt results"
                )
            if suffix not in TEXT_SUFFIXES:
                raise ValueError(
                    f"registered upload file {name!r} has an unaccepted suffix {suffix!r} "
                    f"(accepted: {sorted(TEXT_SUFFIXES)})"
                )
            yield DocumentPayload(
                source_uri=path.resolve().as_uri(),
                raw=path.read_text(encoding="utf-8"),
                mime=TEXT_SUFFIXES[suffix],
                metadata=metadata_by_filename[name],
            )
        return
    sidecars = _collect_sidecars(root, metadata_schema)
    for path in sorted(root.rglob("*")):
        suffix = path.suffix.lower()
        if not path.is_file() or suffix not in TEXT_SUFFIXES:
            continue
        sidecar = sidecars.pop(path.resolve(), None)
        yield DocumentPayload(
            source_uri=path.resolve().as_uri(),
            raw=path.read_text(encoding="utf-8"),
            mime=TEXT_SUFFIXES[suffix],
            metadata=(
                build_envelope(
                    connector="text-directory",
                    original_filename=path.name,
                    context=sidecar.get("context"),
                    governance=sidecar.get("governance"),
                )
                if sidecar is not None
                else {"filename": path.name}
            ),
        )
    if sidecars:
        # an orphan sidecar means someone's metadata is silently NOT applied
        # (a typo'd target name) — that is a config error, not a skippable file
        orphans = sorted(str(p) for p in sidecars)
        raise ValueError(
            f"metadata sidecar(s) without a matching document file: {orphans} — "
            "each <name>.meta.json must sit beside the <name> it describes"
        )


#: MCP10 sidecar convention: a plain-directory source file <name> may carry a
#: display/provenance envelope in <name>.meta.json beside it — the path that
#: finally gives citations an end-user-presentable source (the dev-machine
#: file:// uri identifies the document; context.title / context.attributes
#: hold what a guide agent may SHOW, e.g. the real public page URL — exposed
#: to agents only through the projects.config metadata_exposure allowlist,
#: DR-010). Uploads capture the same envelope through the API instead.
_SIDECAR_SUFFIX = ".meta.json"
_SIDECAR_KEYS = {"context", "governance"}


def _collect_sidecars(
    root: Path, metadata_schema: MetadataSchema | None
) -> dict[Path, dict[str, Any]]:
    """``{target-file-resolved-path: parsed sidecar}`` for every
    ``<name>.meta.json`` under ``root``. A corrupt sidecar is a LOUD failure
    (same rule as an undecodable source file); unknown top-level keys are
    rejected — ``system``/``schema_version`` are server-stamped namespaces a
    sidecar must not forge (DR-010 rule 1/4). Sidecar attributes are validated
    against ``metadata_schema`` (``None`` = the empty schema, fail-closed) —
    see :func:`read_text_documents`."""
    schema = metadata_schema if metadata_schema is not None else MetadataSchema(attributes={})
    collected: dict[Path, dict[str, Any]] = {}
    for sidecar_path in sorted(root.rglob(f"*{_SIDECAR_SUFFIX}")):
        if not sidecar_path.is_file():
            continue
        try:
            parsed = json.loads(sidecar_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            raise ValueError(f"metadata sidecar {sidecar_path} is not valid JSON: {exc}") from exc
        if not isinstance(parsed, dict):
            raise ValueError(f"metadata sidecar {sidecar_path} must be a JSON object")
        unknown = set(parsed) - _SIDECAR_KEYS
        if unknown:
            raise ValueError(
                f"metadata sidecar {sidecar_path} has unknown key(s) {sorted(unknown)} — "
                f"allowed: {sorted(_SIDECAR_KEYS)} (system/schema_version are "
                "server-stamped and cannot be forged from a sidecar)"
            )
        _validate_sidecar_context(sidecar_path, parsed.get("context"))
        try:
            schema.validate_context(parsed.get("context") or {})
        except MetadataValidationError as exc:
            raise ValueError(
                f"metadata sidecar {sidecar_path}: {exc} — sidecar attributes are "
                "validated against the project metadata_schema, the same fence the "
                "upload boundary applies at capture"
            ) from exc
        stem = sidecar_path.name[: -len(_SIDECAR_SUFFIX)]
        if not stem:
            raise ValueError(
                f"metadata sidecar {sidecar_path} names no target file — the "
                "convention is <name>.meta.json beside <name>"
            )
        target = sidecar_path.with_name(stem)
        collected[target.resolve()] = parsed
    return collected


def _validate_sidecar_context(sidecar_path: Path, context: Any) -> None:
    """The sidecar's ``context`` must be the CLOSED core shape — mirroring
    the upload boundary's ``extra=\"forbid\"`` model and the frozen
    ``DocumentMetadataContext`` (``additionalProperties:false``). Without
    this, ``build_envelope``'s normalization SILENTLY drops any other
    context key: the single most likely operator mistake —
    ``context.source_url`` instead of ``context.attributes.source_url`` —
    would quietly yield an envelope with no presentable provenance, the
    exact silent failure MCP10 exists to fix (gate-2 blocker)."""
    if context is None:
        return
    if not isinstance(context, dict):
        raise ValueError(f"metadata sidecar {sidecar_path}: context must be a JSON object")
    unknown = set(context) - {"title", "document_type", "attributes"}
    if unknown:
        raise ValueError(
            f"metadata sidecar {sidecar_path}: unknown context key(s) {sorted(unknown)} — "
            "the closed core is title/document_type/attributes; project fields "
            "(e.g. source_url) belong INSIDE context.attributes"
        )
    for field_name in ("title", "document_type"):
        value = context.get(field_name)
        if value is not None and not isinstance(value, str):
            raise ValueError(
                f"metadata sidecar {sidecar_path}: context.{field_name} must be a "
                f"string or null, got {type(value).__name__}"
            )
    attributes = context.get("attributes")
    if attributes is not None and not isinstance(attributes, dict):
        raise ValueError(
            f"metadata sidecar {sidecar_path}: context.attributes must be a JSON object"
        )


def read_csv_rows(path: Path, *, table: str, pk_column: str) -> Iterator[DocumentPayload]:
    """Yield one payload per CSV row — the structured-source connector (§2/§6).

    Each row's raw text is its canonical JSON (sorted keys, so column order
    changes don't mint new content hashes). ``table`` and the row's pk land
    in metadata because §27.2 row refs cite ``table + pk`` — and the pk must
    exist and be non-empty NOW: a row that cannot be cited later is refused
    at the door, not discovered broken at query time.
    """
    if not table.strip():
        raise ValueError(
            "table name must be non-empty — §27.2 row refs cite table + pk, "
            "and rows ingested under a blank table could never be cited"
        )
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None or pk_column not in reader.fieldnames:
            raise ValueError(
                f"pk column {pk_column!r} not in {path} header "
                f"(found: {list(reader.fieldnames or [])})"
            )
        seen: set[str] = set()
        for index, row in enumerate(reader):
            pk = (row.get(pk_column) or "").strip()
            if not pk:
                raise ValueError(
                    f"row {index} of {path} has an empty pk ({pk_column!r}) — "
                    "it could never be cited as a §27.2 row source ref"
                )
            if pk in seen:
                raise ValueError(
                    f"row {index} of {path} repeats pk {pk!r} — a duplicated "
                    "(table, pk) makes the §27.2 row source ref ambiguous "
                    f"(mis-declared pk column, or dirty data)"
                )
            seen.add(pk)
            yield DocumentPayload(
                source_uri=f"{path.resolve().as_uri()}#{pk_column}={pk}",
                raw=json.dumps(row, ensure_ascii=False, sort_keys=True),
                mime=STRUCTURED_MIME,
                metadata={"table": table, "pk": pk},
            )


#: Stop scanning an xlsx sheet after this many CONSECUTIVE blank rows: real
#: workbooks carry template tails (pre-formatted empty rows), and a sheet's
#: self-reported dimension can lie outright (a pilot file claimed ~1M rows) —
#: trusting either would ingest garbage or spin forever. A content gap this
#: long is a tail, not data (the pilot preprocessor's validated threshold).
XLSX_BLANK_ROW_STOP = 50


def _normalize_header(header: str) -> str:
    """A header cell → its mapping name: trailing parenthetical annotations are
    authoring guidance, not identity (``問題(必填)`` and ``問題`` are the same
    column — half- and full-width parens both appear in real workbooks). The
    full-width pair (U+FF08/U+FF09) is visually identical to half-width in
    most fonts and silently degraded to half-width once already (reviewer
    catch) — the unit test asserts the pattern's actual codepoints, so a
    re-mangled pattern turns the suite red instead of shipping."""
    return re.sub(r"[(（][^()（）]*[)）]\s*$", "", header.strip()).strip()


def _cell_text(value: Any) -> str:
    """One typed openpyxl cell → render text. Floats that carry no fraction are
    the id-column dirt named by the pilot (``編號`` reads back as ``7.0``) —
    canonicalize them to their integer spelling so ids and rendered values
    match what the author typed."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_xlsx_rows(
    path: Path,
    *,
    title_column: str,
    body_column: str,
    id_column: str | None = None,
    extra_columns: Sequence[str] = (),
    label: str | None = None,
) -> Iterator[DocumentPayload]:
    """Yield one per-row TEXT payload from an xlsx workbook's first sheet (SRC1).

    The column MAPPING comes from the caller (source metadata) — this connector
    hard-codes no vocabulary, so any workbook whose rows have "a heading and a
    body" fits (不失一般性; the pilot's two families are just two mappings).
    Each row renders to the pilot-validated text shape::

        【label】title      (or 【title】 when no label)
        id_column:id        (when the id column is mapped and the cell non-blank)
        extra:value         (each non-blank extra column, mapping order)

        body

    The authored id line doubles as ROW IDENTITY: ``content_hash(raw)`` is the
    document identity (§18), so it is what keeps two same-content rows with
    different authored ids from collapsing in ingest dedup (which would drop
    the second row's ``#row=`` provenance). Blank-id rows carry no such line —
    identical blank-id rows are true duplicates and dedup by design (the text
    family's semantics; ingest reports them as skipped).

    which flows into the ordinary text pipeline (chunking + LLM extraction —
    per-row TEXT documents were what the pilot validated, so xlsx documents are
    text-mime and ontology-gated like every other text source).

    Dirt defenses (each one observed in real workbooks):
    - headers carry authoring annotations (``問題(必填)``) → normalized before
      the mapping lookup; a mapped column missing AFTER normalization fails
      loud naming the headers that do exist.
    - the sheet dimension can lie and template tails run long → NO-CONTENT
      rows (blank title AND blank body — fully blank or pre-numbered template
      rows alike) are skipped, and ``XLSX_BLANK_ROW_STOP`` consecutive
      no-content rows end the scan. Pre-numbered tails count toward the stop:
      an id column filled down the template must not keep the scan alive
      (Codex #85).
    - the id column reads back as floats (``7.0``) → canonicalized; a blank id
      falls back to the 1-based data-row ordinal (the pilot's rule: real files
      leave 編號 blank mid-sheet, and refusing them would block whole corpora);
      a DUPLICATE id fails loud — two rows sharing ``#row=`` would make the
      citation ambiguous (mis-declared id column, or dirty data). The ordinal
      counts CONTENT rows only (skipped no-content rows don't consume one),
      so a re-export that inserts/removes rows still SHIFTS blank-id
      citations — the explicit id column is the stable identity; the ordinal
      is a best-effort fallback, not a promise across re-exports.
    """
    for name, value in (("title_column", title_column), ("body_column", body_column)):
        if not value.strip():
            raise ValueError(f"{name} must be non-empty — the render needs it")
    import openpyxl  # heavy import, deliberately local (mirrors the lazy-IO stance)

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        raw_header = next(rows, None)
        if raw_header is None:
            raise ValueError(f"{path} first sheet is empty — no header row to map columns on")
        header = [_normalize_header(_cell_text(h)) for h in raw_header]
        columns: dict[str, int] = {}
        duplicated: set[str] = set()
        for i, h in enumerate(header):
            if not h:
                continue
            if h in columns:
                duplicated.add(h)
            else:
                columns[h] = i
        wanted = [title_column, body_column, *([id_column] if id_column else []), *extra_columns]
        # a MAPPED name that appears twice after normalization (e.g. 問題(必填)
        # and 問題 collapse to the same name) is ambiguous — the mapping only
        # names the normalized header, so it cannot say which column is meant;
        # binding first-wins would silently render the wrong column. Duplicates
        # among UNMAPPED columns stay tolerated: the render never reads them,
        # and refusing the whole workbook for decorative junk would over-block.
        ambiguous = sorted(set(wanted) & duplicated)
        if ambiguous:
            raise ValueError(
                f"mapped column(s) {ambiguous!r} appear more than once in {path} header "
                "after normalization — the mapping cannot say which column is meant; "
                "rename the duplicate columns in the workbook"
            )
        missing = [c for c in wanted if c not in columns]
        if missing:
            raise ValueError(
                f"mapped column(s) {missing!r} not in {path} header after normalization "
                f"(found: {[h for h in header if h]})"
            )

        def cell(row: tuple[Any, ...], column: str) -> str:
            index = columns[column]
            return _cell_text(row[index]) if index < len(row) else ""

        seen: set[str] = set()
        blank_streak = 0
        ordinal = 0
        for row in rows:
            title = cell(row, title_column)
            body = cell(row, body_column)
            if not title and not body:
                # NO CONTENT — whether fully blank or a pre-numbered/
                # pre-formatted template row (id filled down, title/body
                # empty). Both count toward the tail stop: a pre-numbered
                # tail would otherwise reset the streak on every row and the
                # stop would never fire over a huge template tail (Codex #85).
                blank_streak += 1
                if blank_streak >= XLSX_BLANK_ROW_STOP:
                    break
                continue
            blank_streak = 0
            ordinal += 1
            id_text = cell(row, id_column) if id_column else ""
            rid = id_text or str(ordinal)
            if rid in seen:
                raise ValueError(
                    f"row {ordinal} of {path} repeats id {rid!r} — a duplicated row id "
                    "makes the per-row citation ambiguous (mis-declared id column, or "
                    "dirty data)"
                )
            seen.add(rid)
            lines = [f"【{label}】{title}" if label else f"【{title}】"]
            if id_text:
                # the authored row id rides the RENDERED text: content_hash(raw)
                # is THE document identity (§18), so two distinct rows whose
                # title/body/extras render identically would otherwise collapse
                # in ingest dedup and silently drop the second row's #row=
                # provenance (Codex #85) — with the id in the text, distinct
                # authored ids mint distinct documents by construction (the
                # structured family gets this from the pk riding the row JSON).
                # Ordinal-FALLBACK ids are deliberately NOT rendered (the cell
                # was blank — printing an invented number would misquote the
                # sheet), so identical blank-id rows keep the text-family
                # semantics: true duplicates dedup, reported as skipped.
                lines.append(f"{id_column}:{id_text}")
            for extra in extra_columns:
                extra_value = cell(row, extra)
                if extra_value:
                    lines.append(f"{extra}:{extra_value}")
            lines.append("")
            lines.append(body)
            yield DocumentPayload(
                source_uri=f"{path.resolve().as_uri()}#row={rid}",
                raw="\n".join(lines).strip() + "\n",
                mime="text/plain",
                metadata={"filename": path.name, "row": rid},
            )
    finally:
        workbook.close()
